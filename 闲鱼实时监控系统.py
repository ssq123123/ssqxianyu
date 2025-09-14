import customtkinter as ctk
import requests
import time
import hashlib
import threading
import queue
import json
import os
import re
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.utils import get_column_letter
from tkinter import filedialog, messagebox
from datetime import datetime, timedelta
import winsound
import os

try:
    from PIL import Image as PILImage
except ImportError:
    PILImage = None

# 设置CustomTkinter主题
ctk.set_appearance_mode("dark")  # 或 "light"
ctk.set_default_color_theme("blue")

# 常量配置
API_URL = "https://h5api.m.goofish.com/h5/mtop.taobao.idlemtopsearch.pc.search/1.0/"
APP_KEY = "34839810"
USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/134.0.0.0 Safari/537.36 Edg/134.0.0.0"
REQUEST_DELAY = 2.0  # 请求间隔时间（秒）
MONITOR_INTERVAL = 30  # 监控间隔时间（秒）
COOKIE_FILE = "xianyu_cookie.json"
FOUND_ITEMS_FILE = "found_items.json"
HISTORY_FILE = "file_history.json"
MAX_WORKERS = 2  # 监控模式下减少线程数
IMAGE_FOLDER = "xianyu_images"
SUPPORTED_IMAGE_FORMATS = ['jpg', 'jpeg', 'png', 'gif', 'bmp']

class XianyuMonitorSystem:
    def __init__(self):
        self.root = ctk.CTk()
        self.root.title("闲鱼智能监控系统 v4.0 - CustomTkinter版")
        self.root.geometry("900x700")
        self.root.minsize(800, 600)  # 设置最小窗口大小
        
        # 状态变量
        self.is_monitoring = False
        self.monitor_thread = None
        self.found_items = set()
        self.low_price_items = set()
        self.file_history = []
        self.cookie = ""
        self.token = ""
        self.last_check_time = None
        
        # 创建日志队列
        self.log_queue = queue.Queue()
        
        # 确保文件夹存在
        if not os.path.exists(IMAGE_FOLDER):
            os.makedirs(IMAGE_FOLDER)
        
        # 加载数据
        self.load_cookie()
        self.load_found_items()
        self.load_file_history()
        
        # 创建界面
        self.create_widgets()
        
        # 启动日志更新线程
        threading.Thread(target=self.update_log, daemon=True).start()
        
        # 设置窗口关闭事件
        self.root.protocol("WM_DELETE_WINDOW", self.on_closing)
    
    def create_widgets(self):
        # 主容器
        main_container = ctk.CTkFrame(self.root)
        main_container.pack(fill="both", expand=True, padx=10, pady=10)
        
        # 左侧设置面板（使用可滚动框架）
        settings_scroll_frame = ctk.CTkScrollableFrame(main_container)
        settings_scroll_frame.pack(side="left", fill="y", padx=(0, 10))
        settings_scroll_frame.configure(width=350)  # 固定宽度
        
        # 设置面板标题
        ctk.CTkLabel(settings_scroll_frame, text="监控设置", font=ctk.CTkFont(size=16, weight="bold")).pack(pady=10)
        
        # Cookie设置
        ctk.CTkLabel(settings_scroll_frame, text="Cookie:").pack(anchor="w", padx=10)
        self.cookie_entry = ctk.CTkTextbox(settings_scroll_frame, height=60, width=320)
        self.cookie_entry.pack(padx=10, pady=(0, 10), fill="x")
        self.cookie_entry.insert("1.0", self.cookie)
        
        # 关键词设置（支持多个关键词）
        keyword_frame = ctk.CTkFrame(settings_scroll_frame)
        keyword_frame.pack(fill="x", padx=10, pady=(0, 10))
        
        ctk.CTkLabel(keyword_frame, text="监控关键词 (最多5个):", font=ctk.CTkFont(weight="bold")).pack(anchor="w", pady=5)
        
        # 创建5个关键词输入框
        self.keyword_entries = []
        self.keyword_enabled = []
        
        for i in range(5):
            keyword_row_frame = ctk.CTkFrame(keyword_frame)
            keyword_row_frame.pack(fill="x", pady=2)
            
            # 启用复选框
            enabled_var = ctk.BooleanVar(value=True if i == 0 else False)
            self.keyword_enabled.append(enabled_var)
            
            checkbox = ctk.CTkCheckBox(keyword_row_frame, text=f"关键词{i+1}:", variable=enabled_var, width=80)
            checkbox.pack(side="left", padx=5)
            
            # 关键词输入框
            entry = ctk.CTkEntry(keyword_row_frame, placeholder_text=f"输入第{i+1}个关键词")
            entry.pack(side="left", fill="x", expand=True, padx=5)
            self.keyword_entries.append(entry)
            
            # 优先级标签
            priority_label = ctk.CTkLabel(keyword_row_frame, text=f"优先级{i+1}")
            priority_label.pack(side="right", padx=5)
        
        # 价格区间设置
        price_frame = ctk.CTkFrame(settings_scroll_frame)
        price_frame.pack(fill="x", padx=10, pady=(0, 10))
        
        ctk.CTkLabel(price_frame, text="价格区间 (元):").pack(anchor="w", pady=5)
        price_input_frame = ctk.CTkFrame(price_frame)
        price_input_frame.pack(fill="x", pady=5)
        
        self.min_price_entry = ctk.CTkEntry(price_input_frame, width=100, placeholder_text="最低价")
        self.min_price_entry.pack(side="left", padx=5, fill="x", expand=True)
        ctk.CTkLabel(price_input_frame, text=" - ").pack(side="left")
        self.max_price_entry = ctk.CTkEntry(price_input_frame, width=100, placeholder_text="最高价")
        self.max_price_entry.pack(side="left", padx=5, fill="x", expand=True)
        
        # 低价预警设置
        alert_frame = ctk.CTkFrame(settings_scroll_frame)
        alert_frame.pack(fill="x", padx=10, pady=(0, 10))
        
        ctk.CTkLabel(alert_frame, text="低价预警:").pack(anchor="w", pady=5)
        alert_input_frame = ctk.CTkFrame(alert_frame)
        alert_input_frame.pack(fill="x", pady=5)
        
        self.enable_alert = ctk.CTkCheckBox(alert_input_frame, text="启用")
        self.enable_alert.pack(side="left", padx=5)
        self.alert_price_entry = ctk.CTkEntry(alert_input_frame, width=120, placeholder_text="预警价格")
        self.alert_price_entry.pack(side="left", padx=5, fill="x", expand=True)
        
        # 监控间隔设置
        ctk.CTkLabel(settings_scroll_frame, text="监控间隔 (秒):").pack(anchor="w", padx=10)
        self.interval_entry = ctk.CTkEntry(settings_scroll_frame, width=120, placeholder_text="30")
        self.interval_entry.pack(padx=10, pady=(0, 10), fill="x")
        self.interval_entry.insert(0, str(MONITOR_INTERVAL))
        
        # 监控页数设置
        ctk.CTkLabel(settings_scroll_frame, text="监控页数:").pack(anchor="w", padx=10)
        page_frame = ctk.CTkFrame(settings_scroll_frame)
        page_frame.pack(fill="x", padx=10, pady=(0, 10))
        
        self.monitor_pages_entry = ctk.CTkEntry(page_frame, width=80, placeholder_text="3")
        self.monitor_pages_entry.pack(side="left", padx=5)
        self.monitor_pages_entry.insert(0, "3")
        
        ctk.CTkLabel(page_frame, text="页 (每页30个商品)").pack(side="left", padx=5)
        
        # 提示标签
        tip_label = ctk.CTkLabel(page_frame, text="💡 建议1-5页，避免过多请求", font=ctk.CTkFont(size=10))
        tip_label.pack(side="right", padx=5)
        
        # 文件设置
        file_frame = ctk.CTkFrame(settings_scroll_frame)
        file_frame.pack(fill="x", padx=10, pady=(0, 10))
        
        ctk.CTkLabel(file_frame, text="常规结果文件:").pack(anchor="w", pady=5)
        file_input_frame = ctk.CTkFrame(file_frame)
        file_input_frame.pack(fill="x", pady=5)
        
        self.excel_path_entry = ctk.CTkEntry(file_input_frame, width=220, placeholder_text="监控结果.xlsx")
        self.excel_path_entry.pack(side="left", padx=5, fill="x", expand=True)
        
        # 文件浏览和历史记录按钮
        file_btn_frame = ctk.CTkFrame(file_input_frame)
        file_btn_frame.pack(side="right", padx=5)
        
        ctk.CTkButton(file_btn_frame, text="浏览", width=50, command=lambda: self.browse_file("normal")).pack(side="left", padx=2)
        ctk.CTkButton(file_btn_frame, text="📁", width=35, command=lambda: self.show_file_history("normal")).pack(side="right", padx=2)
        
        ctk.CTkLabel(file_frame, text="低价预警文件:").pack(anchor="w", pady=5)
        alert_file_frame = ctk.CTkFrame(file_frame)
        alert_file_frame.pack(fill="x", pady=5)
        
        self.alert_file_entry = ctk.CTkEntry(alert_file_frame, width=220, placeholder_text="低价预警.xlsx")
        self.alert_file_entry.pack(side="left", padx=5, fill="x", expand=True)
        
        # 文件浏览和历史记录按钮
        alert_btn_frame = ctk.CTkFrame(alert_file_frame)
        alert_btn_frame.pack(side="right", padx=5)
        
        ctk.CTkButton(alert_btn_frame, text="浏览", width=50, command=lambda: self.browse_file("alert")).pack(side="left", padx=2)
        ctk.CTkButton(alert_btn_frame, text="📁", width=35, command=lambda: self.show_file_history("alert")).pack(side="right", padx=2)
        
        # 选项设置
        options_frame = ctk.CTkFrame(settings_scroll_frame)
        options_frame.pack(fill="x", padx=10, pady=(0, 10))
        
        self.multi_images_check = ctk.CTkCheckBox(options_frame, text="多图片获取")
        self.multi_images_check.pack(anchor="w", pady=5)
        self.multi_images_check.select()
        
        self.sound_alert_check = ctk.CTkCheckBox(options_frame, text="声音提醒")
        self.sound_alert_check.pack(anchor="w", pady=5)
        self.sound_alert_check.select()
        
        # 控制按钮
        button_frame = ctk.CTkFrame(settings_scroll_frame)
        button_frame.pack(fill="x", padx=10, pady=10)
        
        self.start_btn = ctk.CTkButton(button_frame, text="开始监控", command=self.start_monitoring)
        self.start_btn.pack(fill="x", pady=5)
        
        self.stop_btn = ctk.CTkButton(button_frame, text="停止监控", command=self.stop_monitoring, state="disabled")
        self.stop_btn.pack(fill="x", pady=5)
        
        self.save_cookie_btn = ctk.CTkButton(button_frame, text="保存Cookie", command=self.save_cookie)
        self.save_cookie_btn.pack(fill="x", pady=5)
        
        self.clear_log_btn = ctk.CTkButton(button_frame, text="清除日志", command=self.clear_log)
        self.clear_log_btn.pack(fill="x", pady=5)
        
        self.clear_history_btn = ctk.CTkButton(button_frame, text="清除记录", command=self.clear_found_items)
        self.clear_history_btn.pack(fill="x", pady=5)
        
        # 右侧信息面板
        info_frame = ctk.CTkFrame(main_container)
        info_frame.pack(side="right", fill="both", expand=True)
        
        # 状态显示
        status_frame = ctk.CTkFrame(info_frame)
        status_frame.pack(fill="x", padx=10, pady=10)
        
        # 状态信息使用网格布局
        status_grid_frame = ctk.CTkFrame(status_frame)
        status_grid_frame.pack(fill="x", pady=5)
        
        # 配置网格
        status_grid_frame.grid_columnconfigure(0, weight=1)
        status_grid_frame.grid_columnconfigure(1, weight=1)
        
        self.status_label = ctk.CTkLabel(status_grid_frame, text="状态: 待机中", font=ctk.CTkFont(size=14, weight="bold"))
        self.status_label.grid(row=0, column=0, sticky="w", padx=5, pady=5)
        
        self.count_label = ctk.CTkLabel(status_grid_frame, text="发现商品: 0 | 低价预警: 0")
        self.count_label.grid(row=0, column=1, sticky="e", padx=5, pady=5)
        
        self.last_check_label = ctk.CTkLabel(status_grid_frame, text="上次检查: 从未")
        self.last_check_label.grid(row=1, column=0, columnspan=2, sticky="w", padx=5, pady=5)
        
        # 关键词状态面板
        keyword_status_frame = ctk.CTkFrame(info_frame)
        keyword_status_frame.pack(fill="x", padx=10, pady=(0, 10))
        
        ctk.CTkLabel(keyword_status_frame, text="关键词监控状态", font=ctk.CTkFont(size=14, weight="bold")).pack(pady=5)
        
        # 创建关键词状态标签
        self.keyword_status_labels = []
        for i in range(5):
            status_text = f"关键词{i+1}: 未设置"
            label = ctk.CTkLabel(keyword_status_frame, text=status_text)
            label.pack(anchor="w", padx=5, pady=1)
            self.keyword_status_labels.append(label)
        
        self.update_keyword_status_display()
        
        # 日志显示
        ctk.CTkLabel(info_frame, text="监控日志", font=ctk.CTkFont(size=16, weight="bold")).pack(pady=10)
        
        # 创建滚动文本框
        log_frame = ctk.CTkFrame(info_frame)
        log_frame.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        
        self.log_text = ctk.CTkTextbox(log_frame, height=350, wrap="word")
        self.log_text.pack(fill="both", expand=True, padx=5, pady=5)
    
    def log_message(self, message):
        """添加日志消息"""
        timestamp = datetime.now().strftime("%H:%M:%S")
        self.log_queue.put(f"[{timestamp}] {message}")
    
    def update_log(self):
        """更新日志显示"""
        while True:
            try:
                while not self.log_queue.empty():
                    message = self.log_queue.get_nowait()
                    self.log_text.insert("end", message + "\n")
                    self.log_text.see("end")
                time.sleep(0.1)
            except Exception:
                time.sleep(0.1)
    
    def clear_log(self):
        """清除日志"""
        self.log_text.delete("1.0", "end")
    
    def load_cookie(self):
        """加载Cookie"""
        try:
            if os.path.exists(COOKIE_FILE):
                with open(COOKIE_FILE, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                    self.cookie = data.get('cookie', '')
                    self.log_message("🔑 已加载保存的Cookie")
        except Exception as e:
            self.log_message(f"⚠️ 加载Cookie失败: {e}")
    
    def save_cookie(self):
        """保存Cookie"""
        try:
            cookie_text = self.cookie_entry.get("1.0", "end-1c").strip()
            if not cookie_text:
                messagebox.showwarning("警告", "Cookie不能为空")
                return
            
            self.cookie = cookie_text
            with open(COOKIE_FILE, 'w', encoding='utf-8') as f:
                json.dump({'cookie': self.cookie}, f, ensure_ascii=False, indent=2)
            self.log_message("✅ Cookie保存成功")
        except Exception as e:
            self.log_message(f"❌ 保存Cookie失败: {e}")
    
    def load_found_items(self):
        """加载已发现商品记录"""
        try:
            if os.path.exists(FOUND_ITEMS_FILE):
                with open(FOUND_ITEMS_FILE, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                    self.found_items = set(data.get('items', []))
                    self.low_price_items = set(data.get('low_price_items', []))
                    self.update_count_display()
                    self.log_message(f"📁 已加载历史记录: {len(self.found_items)} 常规 | {len(self.low_price_items)} 低价")
        except Exception as e:
            self.log_message(f"⚠️ 加载历史记录失败: {e}")
    
    def save_found_items(self):
        """保存已发现商品记录"""
        try:
            with open(FOUND_ITEMS_FILE, 'w', encoding='utf-8') as f:
                json.dump({
                    'items': list(self.found_items),
                    'low_price_items': list(self.low_price_items)
                }, f, ensure_ascii=False, indent=2)
        except Exception as e:
            self.log_message(f"⚠️ 保存历史记录失败: {e}")
    
    def clear_found_items(self):
        """清除历史记录"""
        self.found_items.clear()
        self.low_price_items.clear()
        self.save_found_items()
        self.update_count_display()
        self.log_message("🗑️ 已清除所有历史记录")
    
    def load_file_history(self):
        """加载文件历史记录"""
        try:
            if os.path.exists(HISTORY_FILE):
                with open(HISTORY_FILE, 'r', encoding='utf-8') as f:
                    self.file_history = json.load(f)
        except Exception:
            self.file_history = []
    
    def save_file_history(self):
        """保存文件历史记录"""
        try:
            with open(HISTORY_FILE, 'w', encoding='utf-8') as f:
                json.dump(self.file_history, f, ensure_ascii=False, indent=2)
        except Exception:
            pass
    
    def add_to_history(self, file_path, file_type):
        """添加到文件历史记录"""
        now = datetime.now().isoformat()
        # 移除重复项
        self.file_history = [h for h in self.file_history if h['path'] != file_path]
        # 添加新记录
        self.file_history.insert(0, {
            'path': file_path,
            'name': os.path.basename(file_path),
            'type': file_type,
            'time': now
        })
        # 保持最多20条记录
        self.file_history = self.file_history[:20]
        self.save_file_history()
    
    def browse_file(self, file_type):
        """浏览文件 - 带历史记录功能"""
        # 创建文件选择对话框
        file_path = filedialog.asksaveasfilename(
            title="选择Excel文件保存位置",
            defaultextension=".xlsx",
            filetypes=[("Excel文件", "*.xlsx"), ("所有文件", "*.*")]
        )
        
        if file_path:
            if file_type == "normal":
                self.excel_path_entry.delete(0, "end")
                self.excel_path_entry.insert(0, file_path)
            else:
                self.alert_file_entry.delete(0, "end")
                self.alert_file_entry.insert(0, file_path)
            
            self.add_to_history(file_path, file_type)
    
    def show_file_history(self, file_type):
        """显示文件历史记录对话框"""
        if not self.file_history:
            messagebox.showinfo("提示", "暂无文件历史记录")
            return
        
        # 过滤对应类型的文件
        filtered_history = [h for h in self.file_history if h.get('type') == file_type]
        
        if not filtered_history:
            messagebox.showinfo("提示", f"暂无{file_type}类型的文件历史记录")
            return
        
        # 创建历史记录窗口
        history_window = ctk.CTkToplevel(self.root)
        history_window.title("文件历史记录")
        history_window.geometry("600x400")
        history_window.transient(self.root)
        history_window.grab_set()
        
        # 历史记录框架
        history_frame = ctk.CTkScrollableFrame(history_window)
        history_frame.pack(fill="both", expand=True, padx=10, pady=10)
        
        # 标题
        title_text = "常规结果文件" if file_type == "normal" else "低价预警文件"
        ctk.CTkLabel(history_frame, text=f"{title_text}历史记录", font=ctk.CTkFont(size=16, weight="bold")).pack(pady=10)
        
        # 按钮框架
        button_frame = ctk.CTkFrame(history_window)
        button_frame.pack(fill="x", padx=10, pady=(0, 10))
        
        ctk.CTkButton(button_frame, text="清空历史", command=lambda: self.clear_file_history(file_type, history_window)).pack(side="right", padx=5)
        
        # 显示历史记录项
        for i, item in enumerate(filtered_history):
            item_frame = ctk.CTkFrame(history_frame)
            item_frame.pack(fill="x", pady=5, padx=10)
            
            # 文件信息
            info_frame = ctk.CTkFrame(item_frame)
            info_frame.pack(side="left", fill="both", expand=True, padx=5, pady=5)
            
            name_label = ctk.CTkLabel(info_frame, text=f"文件名: {item['name']}", font=ctk.CTkFont(weight="bold"))
            name_label.pack(anchor="w", padx=5)
            
            path_label = ctk.CTkLabel(info_frame, text=f"路径: {item['path']}")
            path_label.pack(anchor="w", padx=5)
            
            time_str = item['time'][:19] if 'T' in item['time'] else item['time']
            time_label = ctk.CTkLabel(info_frame, text=f"使用时间: {time_str}")
            time_label.pack(anchor="w", padx=5)
            
            # 操作按钮
            action_frame = ctk.CTkFrame(item_frame)
            action_frame.pack(side="right", padx=5, pady=5)
            
            select_btn = ctk.CTkButton(action_frame, text="选择", width=60, 
                                     command=lambda path=item['path'], w=history_window: self.select_from_history(path, file_type, w))
            select_btn.pack(pady=2)
            
            delete_btn = ctk.CTkButton(action_frame, text="删除", width=60, fg_color="red", 
                                     command=lambda idx=i, w=history_window: self.delete_history_item(idx, file_type, w))
            delete_btn.pack(pady=2)
    
    def select_from_history(self, file_path, file_type, window):
        """从历史记录中选择文件"""
        if file_type == "normal":
            self.excel_path_entry.delete(0, "end")
            self.excel_path_entry.insert(0, file_path)
        else:
            self.alert_file_entry.delete(0, "end")
            self.alert_file_entry.insert(0, file_path)
        
        window.destroy()
        self.log_message(f"✅ 已选择文件: {os.path.basename(file_path)}")
    
    def delete_history_item(self, index, file_type, window):
        """删除历史记录项"""
        # 过滤对应类型的文件
        filtered_history = [h for h in self.file_history if h.get('type') == file_type]
        
        if index < len(filtered_history):
            item_to_remove = filtered_history[index]
            self.file_history.remove(item_to_remove)
            self.save_file_history()
            window.destroy()
            self.show_file_history(file_type)  # 重新显示
            self.log_message(f"🗑️ 已删除历史记录: {item_to_remove['name']}")
    
    def clear_file_history(self, file_type, window):
        """清空特定类型的文件历史记录"""
        if messagebox.askyesno("确认", f"确定要清空所有{file_type}类型的文件历史记录吗？"):
            self.file_history = [h for h in self.file_history if h.get('type') != file_type]
            self.save_file_history()
            window.destroy()
            self.log_message(f"🗑️ 已清空{file_type}类型的文件历史记录")
    
    def update_keyword_status_display(self):
        """更新关键词状态显示"""
        for i, (entry, enabled_var, label) in enumerate(zip(self.keyword_entries, self.keyword_enabled, self.keyword_status_labels)):
            keyword = entry.get().strip()
            if enabled_var.get() and keyword:
                status_text = f"关键词{i+1}: ✅ {keyword}"
                label.configure(text=status_text, text_color="green")
            elif keyword:
                status_text = f"关键词{i+1}: ⏸️ {keyword} (已禁用)"
                label.configure(text=status_text, text_color="orange")
            else:
                status_text = f"关键词{i+1}: ❌ 未设置"
                label.configure(text=status_text, text_color="gray")
    
    def update_count_display(self):
        """更新计数显示"""
        self.count_label.configure(text=f"发现商品: {len(self.found_items)} | 低价预警: {len(self.low_price_items)}")
    
    def extract_token(self):
        """从cookie中提取token"""
        try:
            cookie = self.cookie_entry.get("1.0", "end-1c").strip()
            if "_m_h5_tk=" not in cookie:
                self.log_message("❌ Cookie中缺少_m_h5_tk值")
                return None
            
            start_idx = cookie.find("_m_h5_tk=") + len("_m_h5_tk=")
            end_idx = cookie.find(";", start_idx)
            if end_idx == -1:
                end_idx = len(cookie)
            
            m_h5_tk_value = cookie[start_idx:end_idx]
            token = m_h5_tk_value.split('_')[0]
            return token
        except Exception as e:
            self.log_message(f"❌ 提取Token失败: {e}")
            return None
    
    def validate_inputs(self):
        """验证输入参数"""
        # 验证Cookie
        self.cookie = self.cookie_entry.get("1.0", "end-1c").strip()
        if not self.cookie:
            messagebox.showwarning("警告", "Cookie不能为空")
            return False
        
        # 提取token
        self.token = self.extract_token()
        if not self.token:
            return False
        
        # 验证关键词
        enabled_keywords = []
        for i, (entry, enabled_var) in enumerate(zip(self.keyword_entries, self.keyword_enabled)):
            if enabled_var.get():
                keyword = entry.get().strip()
                if keyword:
                    enabled_keywords.append(keyword)
                else:
                    messagebox.showwarning("警告", f"关键词{i+1}已启用但为空，请输入关键词或取消启用")
                    return False
        
        if not enabled_keywords:
            messagebox.showwarning("警告", "请至少启用一个关键词")
            return False
        
        # 验证监控间隔
        try:
            interval = int(self.interval_entry.get() or MONITOR_INTERVAL)
            if interval < 10:
                messagebox.showwarning("警告", "监控间隔不能小于10秒")
                return False
        except ValueError:
            messagebox.showwarning("警告", "监控间隔必须是数字")
            return False
        
        # 验证监控页数
        try:
            pages = int(self.monitor_pages_entry.get() or 3)
            if pages < 1 or pages > 10:
                messagebox.showwarning("警告", "监控页数必须在1-10之间")
                return False
        except ValueError:
            messagebox.showwarning("警告", "监控页数必须是数字")
            return False
        
        return True
    
    def start_monitoring(self):
        """开始智能监控"""
        if self.is_monitoring:
            return
        
        if not self.validate_inputs():
            return
        
        self.is_monitoring = True
        self.start_btn.configure(state="disabled")
        self.stop_btn.configure(state="normal")
        self.status_label.configure(text="状态: 监控中...")
        
        # 启动监控线程
        self.monitor_thread = threading.Thread(target=self.intelligent_monitor_loop, daemon=True)
        self.monitor_thread.start()
        
        # 获取启用的关键词
        enabled_keywords = []
        for i, (entry, enabled_var) in enumerate(zip(self.keyword_entries, self.keyword_enabled)):
            if enabled_var.get():
                keyword = entry.get().strip()
                if keyword:
                    enabled_keywords.append(keyword)
        
        interval = int(self.interval_entry.get() or MONITOR_INTERVAL)
        monitor_pages = int(self.monitor_pages_entry.get() or 3)
        
        self.log_message(f"🚀 开始智能多关键词监控:")
        for i, keyword in enumerate(enabled_keywords, 1):
            self.log_message(f"  {i}. {keyword}")
        self.log_message(f"⏰ 监控间隔: {interval} 秒")
        self.log_message(f"📄 每个关键词监控: {monitor_pages} 页 (最多{monitor_pages * 30}个商品)")
        
        if self.enable_alert.get():
            alert_price = self.alert_price_entry.get().strip()
            self.log_message(f"🚨 低价预警已启用: 低于 {alert_price} 元")
        
        self.log_message(f"🔍 开始多关键词监控，覆盖更全面的商品信息...")
    
    def stop_monitoring(self):
        """停止监控"""
        self.is_monitoring = False
        self.start_btn.configure(state="normal")
        self.stop_btn.configure(state="disabled")
        self.status_label.configure(text="状态: 待机中")
        self.log_message("⏹️ 监控已停止")
    
    def intelligent_monitor_loop(self):
        """智能监控循环 - 支持多关键词多页搜索"""
        # 获取启用的关键词
        enabled_keywords = []
        for entry, enabled_var in zip(self.keyword_entries, self.keyword_enabled):
            if enabled_var.get():
                keyword = entry.get().strip()
                if keyword:
                    enabled_keywords.append(keyword)
        
        interval = int(self.interval_entry.get() or MONITOR_INTERVAL)
        monitor_pages = int(self.monitor_pages_entry.get() or 3)
        
        while self.is_monitoring:
            try:
                self.log_message(f"🔍 开始新一轮监控 - 共{len(enabled_keywords)}个关键词")
                
                # 统计所有关键词的结果
                total_new_items = []
                total_low_price_items = []
                total_products_checked = 0
                
                # 遍历所有启用的关键词
                for keyword_index, keyword in enumerate(enabled_keywords, 1):
                    if not self.is_monitoring:
                        break
                    
                    self.log_message(f"📝 [{keyword_index}/{len(enabled_keywords)}] 正在搜索: '{keyword}'")
                    
                    keyword_new_items = []
                    keyword_low_price_items = []
                    keyword_products_checked = 0
                    
                    # 搜索当前关键词的多页商品
                    for page in range(1, monitor_pages + 1):
                        if not self.is_monitoring:
                            break
                        
                        self.log_message(f"   📄 [{keyword}] 第{page}页")
                        products = self.fetch_products_by_page(keyword, page)
                        
                        if products:
                            keyword_products_checked += len(products)
                            
                            for product in products:
                                parsed = self.parse_product(product)
                                if not parsed:
                                    continue
                                
                                item_id = parsed['item_id']
                                
                                # 检查是否是新商品（基于ID去重）
                                if item_id in self.found_items or item_id in self.low_price_items:
                                    continue
                                
                                # 添加关键词标记到商品信息中
                                parsed['source_keyword'] = keyword
                                parsed['keyword_index'] = keyword_index
                                
                                price = self.parse_price(parsed['price'])
                                
                                # 检查低价预警
                                if self.enable_alert.get():
                                    try:
                                        alert_price = float(self.alert_price_entry.get() or 999999)
                                        if price > 0 and price <= alert_price:
                                            keyword_low_price_items.append(parsed)
                                            self.low_price_items.add(item_id)
                                            continue
                                    except ValueError:
                                        pass
                                
                                # 检查价格区间
                                if self.is_price_in_range(parsed['price']):
                                    keyword_new_items.append(parsed)
                                    self.found_items.add(item_id)
                        else:
                            self.log_message(f"   ⚠️ [{keyword}] 第{page}页搜索失败")
                        
                        # 页间延迟
                        if page < monitor_pages and self.is_monitoring:
                            time.sleep(0.5)
                    
                    # 统计当前关键词的结果
                    total_new_items.extend(keyword_new_items)
                    total_low_price_items.extend(keyword_low_price_items)
                    total_products_checked += keyword_products_checked
                    
                    self.log_message(f"   📊 [{keyword}] 检查{keyword_products_checked}个商品, 发现{len(keyword_new_items)}个新商品, {len(keyword_low_price_items)}个低价预警")
                    
                    # 关键词间延迟，避免请求过频
                    if keyword_index < len(enabled_keywords) and self.is_monitoring:
                        time.sleep(1)
                
                # 统计本轮监控的总结果
                self.log_message(f"📈 本轮监控总计: 检查{total_products_checked}个商品, 发现{len(total_new_items)}个新商品, {len(total_low_price_items)}个低价预警")
                    
                # 处理发现的商品
                if total_new_items or total_low_price_items:
                    if total_new_items:
                        self.save_new_items(total_new_items, "normal")
                        for item in total_new_items:
                            price = self.parse_price(item['price'])
                            keyword_info = f"[{item['source_keyword']}]"
                            self.log_message(f"📦 {keyword_info} 发现新商品: {item['description'][:30]}... 价格: {price}元")
                    
                    if total_low_price_items:
                        self.save_new_items(total_low_price_items, "alert")
                        for item in total_low_price_items:
                            price = self.parse_price(item['price'])
                            keyword_info = f"[{item['source_keyword']}]"
                            self.log_message(f"🚨 {keyword_info} 低价预警: {item['description'][:30]}... 价格: {price}元")
                    
                    self.save_found_items()
                    self.update_count_display()
                    
                    # 发出提醒
                    total_new = len(total_new_items) + len(total_low_price_items)
                    self.show_notification(total_new, len(total_low_price_items))
                else:
                    self.log_message("✅ 未发现新商品")
                
                # 更新最后检查时间
                self.last_check_time = datetime.now()
                self.last_check_label.configure(text=f"上次检查: {self.last_check_time.strftime('%H:%M:%S')}")
                
                # 等待下次检查
                for i in range(interval):
                    if not self.is_monitoring:
                        break
                    time.sleep(1)
                    remaining = interval - i - 1
                    if remaining > 0:
                        self.status_label.configure(text=f"状态: 监控中... {remaining}秒后下次检查")
                
            except Exception as e:
                self.log_message(f"⚠️ 监控过程出错: {str(e)}")
                time.sleep(10)
    
    def fetch_products_by_page(self, keyword, page=1):
        """按页获取商品（按时间排序）"""
        try:
            # 生成签名和请求参数 - 按时间排序
            sign, timestamp, request_data = self.generate_sign_with_sort_and_page(keyword, page)
            
            headers = {
                "cookie": self.cookie,
                "origin": "https://www.goofish.com",
                "referer": "https://www.goofish.com/",
                "user-agent": USER_AGENT
            }
            
            params = {
                "jsv": "2.7.2",
                "appKey": APP_KEY,
                "t": timestamp,
                "sign": sign,
                "v": "1.0",
                "type": "originaljson",
                "accountSite": "xianyu",
                "dataType": "json",
                "timeout": "20000",
                "api": "mtop.taobao.idlemtopsearch.pc.search"
            }
            
            response = requests.post(
                url=API_URL,
                headers=headers,
                params=params,
                data={"data": request_data},
                timeout=15
            )
            
            response.raise_for_status()
            result = response.json()
            
            # 检查Token是否过期
            if "ret" in result and "FAIL_SYS_TOKEN_EXOIRED" in str(result["ret"]):
                self.log_message("❌ Token已过期，请更新Cookie")
                return None
            
            if "data" in result and "resultList" in result["data"]:
                return result["data"]["resultList"]
            else:
                self.log_message(f"⚠️ 第{page}页响应数据格式异常")
                return None
        
        except Exception as e:
            self.log_message(f"❌ 第{page}页请求失败: {str(e)}")
            return None
    
    def generate_sign_with_sort_and_page(self, keyword, page):
        """生成带排序和页数的签名 - 按发布时间排序"""
        timestamp = int(time.time() * 1000)
        
        # 构建请求数据 - 添加时间排序和页数
        request_data = (
            f'{{"pageNumber":{page},"keyword":"{keyword}","fromFilter":false,'
            f'"rowsPerPage":30,"sortValue":"gmtCreate","sortField":"DESC","customDistance":"",'
            f'"gps":"","propValueStr":"","customGps":"","searchReqFromPage":"pcSearch",'
            f'"extraFilterValue":"","userPositionJson":""}}'
        )
        
        # 构建签名原始字符串
        sign_str = f"{self.token}&{timestamp}&{APP_KEY}&{request_data}"
        
        # 计算MD5签名
        md5 = hashlib.md5()
        md5.update(sign_str.encode("utf-8"))
        sign = md5.hexdigest()
        
        return sign, timestamp, request_data
    
    def parse_product(self, product):
        """解析商品数据"""
        try:
            item_data = product["data"]["item"]["main"]["exContent"]
            click_params = product["data"]["item"]["main"]["clickParam"]["args"]
            
            # 提取基本信息
            pic_url = item_data.get("picUrl", "") or click_params.get("picUrl", "无图片链接")
            user_name = item_data.get("userNick", "未知用户").strip()
            title = item_data.get("title", "").strip()
            post_fee = click_params.get("tagname", "不包邮")
            description = f"{post_fee} +++ {title}"
            item_id = item_data.get("itemId", "")
            product_url = f"https://www.goofish.com/item?id={item_id}"
            price = click_params.get("price", "未知")
            area = item_data.get("area", "未知地区").strip()
            
            # 提取多图片
            all_images = [pic_url] if pic_url != "无图片链接" else []
            
            if self.multi_images_check.get():
                try:
                    full_item = product["data"]["item"]
                    found_urls = self.extract_image_urls(full_item)
                    for url in found_urls:
                        if url and url not in all_images and ('http' in url or '//' in url):
                            if not any(exclude in url.lower() for exclude in ['avatar', 'icon', 'logo']):
                                all_images.append(url)
                except Exception:
                    pass
            
            return {
                "user_name": user_name,
                "description": description,
                "url": product_url,
                "price": price,
                "area": area,
                "pic_url": pic_url,
                "item_id": item_id,
                "all_images": all_images
            }
        
        except Exception as e:
            self.log_message(f"⚠️ 商品解析失败: {str(e)}")
            return None
    
    def extract_image_urls(self, data, found_urls=None):
        """递归提取图片URL"""
        if found_urls is None:
            found_urls = set()
        
        if isinstance(data, dict):
            for key, value in data.items():
                if any(keyword in key.lower() for keyword in ['pic', 'img', 'image', 'photo']):
                    if isinstance(value, str) and ('http' in value or '//' in value):
                        found_urls.add(value)
                    elif isinstance(value, list):
                        for item in value:
                            if isinstance(item, str) and ('http' in item or '//' in item):
                                found_urls.add(item)
                            elif isinstance(item, dict) and 'url' in item:
                                found_urls.add(item['url'])
                else:
                    self.extract_image_urls(value, found_urls)
        elif isinstance(data, list):
            for item in data:
                self.extract_image_urls(item, found_urls)
        
        return found_urls
    
    def parse_price(self, price_str):
        """解析价格"""
        try:
            price_clean = re.sub(r'[^\d.]', '', str(price_str))
            return float(price_clean) if price_clean else 0
        except:
            return 0
    
    def is_price_in_range(self, price_str):
        """检查价格是否在范围内"""
        try:
            price = self.parse_price(price_str)
            min_price = float(self.min_price_entry.get() or 0)
            max_price = float(self.max_price_entry.get() or 999999)
            return min_price <= price <= max_price
        except:
            return True
    
    def download_image(self, pic_url, item_id, img_index=0):
        """下载图片"""
        try:
            if pic_url == "无图片链接":
                return None
            
            if not pic_url.startswith(('http://', 'https://')):
                pic_url = f"http:{pic_url}" if pic_url.startswith('//') else f"https://{pic_url}"
            
            file_ext = pic_url.split(".")[-1].split("?")[0].lower()
            if file_ext not in SUPPORTED_IMAGE_FORMATS:
                file_ext = "jpg"
            
            if img_index == 0:
                file_name = f"{IMAGE_FOLDER}/{item_id}.{file_ext}"
            else:
                file_name = f"{IMAGE_FOLDER}/{item_id}_{img_index}.{file_ext}"
            
            if os.path.exists(file_name):
                return file_name
            
            headers = {"User-Agent": USER_AGENT}
            response = requests.get(pic_url, headers=headers, timeout=10)
            response.raise_for_status()
            
            with open(file_name, "wb") as f:
                f.write(response.content)
            
            return file_name
        
        except Exception:
            return None
    
    def save_new_items(self, items, item_type):
        """保存新商品到Excel"""
        try:
            # 确定文件路径
            if item_type == "alert":
                excel_path = self.alert_file_entry.get().strip() or "低价预警.xlsx"
            else:
                excel_path = self.excel_path_entry.get().strip() or "监控结果.xlsx"
            
            # 处理Excel文件
            if os.path.exists(excel_path):
                wb = load_workbook(excel_path)
                ws = wb.active
                next_row = ws.max_row + 1
            else:
                wb = Workbook()
                ws = wb.active
                headers = ["发现时间", "类型", "关键词", "用户名", "商品描述", "价格", "地区", "商品链接"]
                if self.multi_images_check.get():
                    headers.extend(["图片1", "图片2", "图片3"])
                ws.append(headers)
                next_row = 2
                
                # 设置列宽
                ws.column_dimensions["A"].width = 20
                ws.column_dimensions["B"].width = 10
                ws.column_dimensions["C"].width = 15
                ws.column_dimensions["D"].width = 15
                ws.column_dimensions["E"].width = 50
                ws.column_dimensions["F"].width = 10
                ws.column_dimensions["G"].width = 15
                ws.column_dimensions["H"].width = 30
            
            # 添加数据
            for item in items:
                current_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                item_type_str = "🚨低价预警" if item_type == "alert" else "📦常规发现"
                
                row_data = [
                    current_time,
                    item_type_str,
                    item.get('source_keyword', '未知'),
                    item["user_name"],
                    item["description"],
                    item["price"],
                    item["area"],
                    item["url"]
                ]
                
                # 处理图片
                if self.multi_images_check.get():
                    all_images = item.get("all_images", [])
                    downloaded_images = []
                    
                    for img_idx, pic_url in enumerate(all_images[:3]):
                        pic_path = self.download_image(pic_url, item["item_id"], img_idx)
                        if pic_path and os.path.exists(pic_path):
                            downloaded_images.append(pic_path)
                    
                    # 插入图片
                    for img_idx, pic_path in enumerate(downloaded_images):
                        try:
                            img = Image(pic_path)
                            img.width = 80
                            img.height = 80
                            col = get_column_letter(9 + img_idx)
                            ws.add_image(img, anchor=f"{col}{next_row}")
                        except Exception:
                            pass
                    
                    if downloaded_images:
                        ws.row_dimensions[next_row].height = 65
                
                ws.append(row_data)
                next_row += 1
            
            wb.save(excel_path)
            
            file_type_str = "低价预警" if item_type == "alert" else "常规监控"
            self.log_message(f"💾 已保存 {len(items)} 条{file_type_str}记录到 {os.path.basename(excel_path)}")
            
        except Exception as e:
            self.log_message(f"❌ 保存Excel失败: {str(e)}")
    
    def show_notification(self, total_count, low_price_count):
        """显示通知"""
        try:
            # 声音提醒
            if self.sound_alert_check.get():
                if low_price_count > 0:
                    # 低价预警使用更急促的声音
                    winsound.Beep(1000, 500)  # 高频短促
                    winsound.Beep(1000, 500)
                else:
                    winsound.Beep(800, 300)  # 低频提醒
            
            # 窗口提醒
            self.root.attributes('-topmost', True)
            self.root.after(2000, lambda: self.root.attributes('-topmost', False))
            
            # 系统通知
            if low_price_count > 0:
                message = f"🚨 发现 {low_price_count} 件低价预警商品！"
            else:
                message = f"📦 发现 {total_count} 件新商品"
            
            self.log_message(f"🔔 {message}")
            
        except Exception:
            pass
    
    def on_closing(self):
        """窗口关闭处理"""
        if self.is_monitoring:
            if messagebox.askokcancel("退出", "监控正在运行，确定要退出吗？"):
                self.stop_monitoring()
                self.root.destroy()
        else:
            self.root.destroy()

def main():
    app = XianyuMonitorSystem()
    app.root.mainloop()

if __name__ == "__main__":
    main()